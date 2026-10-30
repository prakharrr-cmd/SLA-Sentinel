"""File parser supporting CSV (with UTF-8 / BOM / delimiters) and Excel (.xlsx) workbooks."""

import csv
import io
import re
from typing import Any, Dict, List, Optional, Tuple
import openpyxl

MAX_FILE_SIZE_BYTES = 50 * 1024 * 1024  # 50 MB
MAX_ROW_COUNT = 100_000
PREVIEW_SAMPLE_LIMIT = 10

SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}


class FileParsingError(Exception):
    """Raised when an uploaded file cannot be parsed or violates format constraints."""
    pass


def infer_data_type(values: List[Any]) -> str:
    """Infer high-level column type from a sample of non-empty values."""
    non_empty = [v for v in values if v is not None and str(v).strip() != ""]
    if not non_empty:
        return "string"

    # Check for boolean
    bool_count = sum(1 for v in non_empty if str(v).strip().lower() in {"true", "false", "yes", "no", "1", "0"})
    if bool_count == len(non_empty):
        return "boolean"

    # Check for integer
    int_count = sum(1 for v in non_empty if re.fullmatch(r"[-+]?\d+", str(v).strip()))
    if int_count == len(non_empty):
        return "integer"

    # Check for float
    float_count = sum(1 for v in non_empty if re.fullmatch(r"[-+]?\d*\.\d+", str(v).strip()))
    if (int_count + float_count) == len(non_empty):
        return "float"

    # Check for date / timestamp
    date_patterns = [
        r"\d{1,2}[-/\.]\d{1,2}[-/\.]\d{2,4}",  # 31-05-2023 or 2023-05-31
        r"\d{4}[-/\.]\d{1,2}[-/\.]\d{1,2}",
        r"\d{1,2}[-/\.]\d{1,2}[-/\.]\d{2,4}\s+\d{1,2}[\.:]\d{2}",  # 31-05-2023 19.14
    ]
    date_count = sum(
        1 for v in non_empty if any(re.search(pat, str(v).strip()) for pat in date_patterns)
    )
    if date_count >= len(non_empty) * 0.7:
        return "datetime"

    return "string"


def parse_csv_stream(
    stream: io.TextIOBase,
    filename: str = "dataset.csv",
    max_rows: int = MAX_ROW_COUNT,
) -> Tuple[List[str], List[Dict[str, Any]], Dict[str, str], List[str], str]:
    """Parse a CSV text stream, auto-detecting delimiters and multi-line fields.

    Returns:
        (columns, rows, inferred_types, warnings, delimiter)
    """
    warnings: List[str] = []

    # Read initial chunk for dialect sniffing
    initial_sample = stream.read(8192)
    if not initial_sample.strip():
        raise FileParsingError("The CSV file is empty.")

    stream.seek(0)

    # Detect delimiter reliably
    delimiter = ","
    try:
        sniffer = csv.Sniffer()
        dialect = sniffer.sniff(initial_sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except Exception:
        # Fallback: Count frequency of delimiter candidates across first non-empty line
        first_line = ""
        for line in initial_sample.splitlines():
            if line.strip():
                first_line = line
                break
        candidate_counts = {delim: first_line.count(delim) for delim in [",", ";", "\t", "|"]}
        best_candidate = max(candidate_counts, key=candidate_counts.get)
        if candidate_counts[best_candidate] > 0:
            delimiter = best_candidate

    reader = csv.reader(stream, delimiter=delimiter)

    # Parse headers
    try:
        header_row = next(reader)
    except StopIteration:
        raise FileParsingError("The CSV file contains no data rows or headers.")

    # Strip BOM and whitespace from header names
    columns = [
        str(col).strip().replace("\ufeff", "") for col in header_row if str(col).strip() != ""
    ]
    if not columns:
        raise FileParsingError("CSV contains no valid column headers.")

    # Check for duplicate column names in source
    seen_cols = set()
    cleaned_columns: List[str] = []
    for idx, col in enumerate(columns):
        if col in seen_cols:
            dedup_name = f"{col}_{idx + 1}"
            warnings.append(f"Duplicate column header '{col}' was renamed to '{dedup_name}'.")
            cleaned_columns.append(dedup_name)
            seen_cols.add(dedup_name)
        else:
            cleaned_columns.append(col)
            seen_cols.add(col)

    columns = cleaned_columns
    num_cols = len(columns)

    rows: List[Dict[str, Any]] = []
    row_count = 0

    for line_num, row_vals in enumerate(reader, start=2):
        if not any(v.strip() for v in row_vals):
            continue  # skip entirely blank lines

        row_count += 1
        if row_count > max_rows:
            raise FileParsingError(
                f"File exceeds maximum allowed limit of {max_rows:,} rows."
            )

        # Pad or trim row_vals to align with columns
        if len(row_vals) < num_cols:
            row_vals = row_vals + [""] * (num_cols - len(row_vals))
        elif len(row_vals) > num_cols:
            row_vals = row_vals[:num_cols]

        row_dict = {col: row_vals[i] for i, col in enumerate(columns)}
        rows.append(row_dict)

    if not rows:
        raise FileParsingError("The CSV file contains column headers but zero data rows.")

    # Type inference based on sample
    sample_for_inference = rows[:50]
    inferred_types: Dict[str, str] = {}
    for col in columns:
        col_values = [r[col] for r in sample_for_inference]
        inferred_types[col] = infer_data_type(col_values)

    return columns, rows, inferred_types, warnings, delimiter


def parse_excel_bytes(
    file_bytes: bytes,
    filename: str = "dataset.xlsx",
    sheet_name: Optional[str] = None,
    max_rows: int = MAX_ROW_COUNT,
) -> Tuple[List[str], List[Dict[str, Any]], Dict[str, str], List[str], List[str], str]:
    """Parse an Excel (.xlsx) workbook bytes.

    Returns:
        (columns, rows, inferred_types, warnings, sheet_names, selected_sheet)
    """
    warnings: List[str] = []

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise FileParsingError(f"Failed to read Excel workbook: {exc}")

    sheet_names = workbook.sheetnames
    if not sheet_names:
        raise FileParsingError("Excel workbook contains no sheets.")

    selected_sheet = sheet_name if (sheet_name and sheet_name in sheet_names) else sheet_names[0]
    sheet = workbook[selected_sheet]

    rows_iter = sheet.iter_rows(values_only=True)

    # First non-empty row as header
    header_row = None
    for r in rows_iter:
        if r and any(cell is not None and str(cell).strip() != "" for cell in r):
            header_row = r
            break

    if not header_row:
        raise FileParsingError(f"Worksheet '{selected_sheet}' contains no headers or data.")

    columns = [
        str(cell).strip() for cell in header_row if cell is not None and str(cell).strip() != ""
    ]
    if not columns:
        raise FileParsingError(f"Worksheet '{selected_sheet}' contains no column headers.")

    # Deduplicate headers
    seen_cols = set()
    cleaned_columns: List[str] = []
    for idx, col in enumerate(columns):
        if col in seen_cols:
            dedup_name = f"{col}_{idx + 1}"
            warnings.append(f"Duplicate column header '{col}' was renamed to '{dedup_name}'.")
            cleaned_columns.append(dedup_name)
            seen_cols.add(dedup_name)
        else:
            cleaned_columns.append(col)
            seen_cols.add(col)
    columns = cleaned_columns
    num_cols = len(columns)

    rows: List[Dict[str, Any]] = []
    row_count = 0

    for r in rows_iter:
        if not r or not any(cell is not None and str(cell).strip() != "" for cell in r):
            continue

        row_count += 1
        if row_count > max_rows:
            raise FileParsingError(
                f"Worksheet '{selected_sheet}' exceeds maximum allowed limit of {max_rows:,} rows."
            )

        row_vals = list(r[:num_cols])
        if len(row_vals) < num_cols:
            row_vals += [None] * (num_cols - len(row_vals))

        row_dict = {
            col: (str(val).strip() if val is not None else "")
            for col, val in zip(columns, row_vals)
        }
        rows.append(row_dict)

    if not rows:
        raise FileParsingError(f"Worksheet '{selected_sheet}' contains headers but zero data rows.")

    sample_for_inference = rows[:50]
    inferred_types: Dict[str, str] = {}
    for col in columns:
        col_values = [r[col] for r in sample_for_inference]
        inferred_types[col] = infer_data_type(col_values)

    return columns, rows, inferred_types, warnings, sheet_names, selected_sheet


def parse_raw_content(
    content: bytes,
    filename: str = "dataset.csv",
    sheet_name: Optional[str] = None,
) -> Dict[str, Any]:
    """Inspect and route file parsing for CSV or Excel files."""
    if not content or len(content) == 0:
        raise FileParsingError("The uploaded file is empty (0 bytes).")

    if len(content) > MAX_FILE_SIZE_BYTES:
        raise FileParsingError(
            f"File size ({len(content) / (1024 * 1024):.2f} MB) exceeds maximum allowed limit of 50 MB."
        )

    lower_name = filename.lower().strip()

    if lower_name.endswith(".xls") and not lower_name.endswith(".xlsx"):
        raise FileParsingError(
            "Legacy Excel .xls format is unsupported. Please convert or export your spreadsheet to .xlsx or .csv format."
        )

    if lower_name.endswith(".xlsx"):
        cols, rows, inf_types, warns, sheets, sel_sheet = parse_excel_bytes(
            content, filename=filename, sheet_name=sheet_name
        )
        file_format = "xlsx"
    elif lower_name.endswith(".csv") or lower_name.endswith(".txt") or filename == "pasted_content.csv":
        # Handle UTF-8 and UTF-8 BOM
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception as e:
                raise FileParsingError(f"Unable to decode CSV text: {e}")

        stream = io.StringIO(text)
        cols, rows, inf_types, warns, detected_delim = parse_csv_stream(stream, filename=filename)
        file_format = "csv"
        sheets = None
        sel_sheet = None
        delimiter = detected_delim
    else:
        raise FileParsingError(
            f"Unsupported file type for '{filename}'. Supported formats are: .csv and .xlsx."
        )

    return {
        "filename": filename,
        "file_format": file_format,
        "detected_columns": cols,
        "rows": rows,
        "total_rows": len(rows),
        "inferred_types": inf_types,
        "warnings": warns,
        "sheet_names": sheets,
        "selected_sheet": sel_sheet,
        "delimiter": delimiter if file_format == "csv" else None,
    }

